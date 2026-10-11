"""Turn uploaded resume files into Claude message content blocks."""

import base64
import io

from docx import Document


def resume_to_content(filename: str, data: bytes) -> dict:
    """Return a content block for one resume.

    PDFs are sent to Claude natively (it reads layout, tables and columns),
    DOCX/TXT files are converted to plain text first.
    """
    name = filename.lower()
    if name.endswith(".pdf"):
        return {
            "type": "document",
            "source": {
                "type": "base64",
                "media_type": "application/pdf",
                "data": base64.standard_b64encode(data).decode("ascii"),
            },
        }
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(p for p in parts if p.strip())
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported file type: {filename}")

    if not text.strip():
        raise ValueError(f"No text could be extracted from {filename}")
    return {"type": "text", "text": f"<resume>\n{text}\n</resume>"}
